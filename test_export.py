import streamlit as st
import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

st.set_page_config(layout="wide")

# 1. THE MOCK DATA
dummy_ai_response = """
### PROFESSIONAL SUMMARY
Recent Psychology graduate and researcher with recognized expertise in legal frameworks, policy analysis, and spatial data mapping.

### EDUCATION & RESEARCH
**Student Researcher** | University | Campus | Spring 2026
- Designed a 2x2 factorial study to analyze juror perceptions and camera perspective bias, utilizing a security guard variable in place of a police officer.

**Awardee** | Academic Department | N/A | 2026
- Awarded the Valeria Dean Burgess Stevens Prize for academic writing and policy research.

### TECHNICAL PROJECTS
**Independent Researcher** | Remote | May 2026
- Mapped Institutional Risk and surveillance landscapes at the county level using Geographic Information Systems (GIS).
"""

# 2. SYNCHRONIZED SETTINGS 
density_settings = {
    1: {"margin_in": 0.5,  "font_pt": 10,   "space_after_pt": 0, "line_height": 1.0},
    2: {"margin_in": 0.6,  "font_pt": 10.5, "space_after_pt": 2, "line_height": 1.1},
    3: {"margin_in": 0.75, "font_pt": 11,   "space_after_pt": 4, "line_height": 1.15},
    4: {"margin_in": 0.85, "font_pt": 11.5, "space_after_pt": 6, "line_height": 1.25},
    5: {"margin_in": 1.0,  "font_pt": 12,   "space_after_pt": 8, "line_height": 1.5}
}

font_settings = {
    "Calibri (Modern & Clean)": {"word_name": "Calibri", "css_family": "'Calibri', sans-serif"},
    "Garamond (Classic & Academic)": {"word_name": "Garamond", "css_family": "'Garamond', serif"},
    "Arial (Standard & Safe)": {"word_name": "Arial", "css_family": "'Arial', sans-serif"},
    "Georgia (Elegant & Readable)": {"word_name": "Georgia", "css_family": "'Georgia', serif"},
    "Times New Roman (Traditional)": {"word_name": "Times New Roman", "css_family": "'Times New Roman', serif"}
}

def lock_paragraph_spacing(paragraph, space_after, space_before=0, line_ht=1.15):
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_ht

# 3. EXPORT FUNCTION
def create_polished_word_doc(resume_text, level=3, font_choice="Calibri (Modern & Clean)"):
    doc = Document()
    settings = density_settings[level]
    selected_font = font_settings[font_choice]["word_name"]
    
    margin = settings["margin_in"]
    font_size = settings["font_pt"]
    spacing = settings["space_after_pt"]
    line_ht = settings["line_height"]

    for section in doc.sections:
        section.top_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(margin)
        section.right_margin = Inches(margin)

    # BUGFIX: Pull the tab stop back by 0.05 inches to prevent wrapping
    right_tab_position = 8.5 - (margin * 2) - 0.05

    style = doc.styles['Normal']
    style.font.name = selected_font
    style.font.size = Pt(font_size)

    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("ALEX")
    header_run.bold = True
    header_run.font.size = Pt(font_size + 6)
    lock_paragraph_spacing(header, space_after=0, line_ht=line_ht)
    
    contact = doc.add_paragraph("Boston, MA | alex@example.com | 555-0199")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lock_paragraph_spacing(contact, space_after=12, line_ht=line_ht)

    for line in resume_text.split('\n'):
        if not line.strip(): continue 
            
        if line.strip().startswith('###'): 
            clean_header = line.replace('###', '').strip().upper()
            p = doc.add_paragraph()
            run = p.add_run(clean_header)
            run.bold = True
            run.font.size = Pt(font_size + 1)
            lock_paragraph_spacing(p, space_after=2, space_before=12, line_ht=line_ht)
            
        elif line.strip().startswith('**'):
            # BUGFIX: Safely handle missing dates or locations
            parts = [p.strip() for p in line.replace('**', '').split('|')]
            p = doc.add_paragraph()
            
            tab_stops = p.paragraph_format.tab_stops
            tab_stops.add_tab_stop(Inches(right_tab_position), WD_TAB_ALIGNMENT.RIGHT)
            
            if len(parts) >= 2:
                p.add_run(f"{parts[0]} | {parts[1]}").bold = True
            
            if len(parts) == 3:
                p.add_run(f"\t{parts[2]}").bold = False
            elif len(parts) >= 4:
                right_text = f"\t{parts[2]} | {parts[3]}".replace("N/A | ", "").replace(" | N/A", "")
                p.add_run(right_text).bold = False
                
            lock_paragraph_spacing(p, space_after=2, space_before=6, line_ht=line_ht)

        elif line.strip().startswith('-'): 
            clean_bullet = line.replace('-', '', 1).strip()
            p = doc.add_paragraph(clean_bullet, style='List Bullet')
            lock_paragraph_spacing(p, space_after=spacing, line_ht=line_ht)
            
        else:
            p = doc.add_paragraph(line.strip())
            lock_paragraph_spacing(p, space_after=spacing, line_ht=line_ht)
            
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# 4. LIVE HTML PREVIEW
def render_live_preview(resume_text, level, font_choice):
    settings = density_settings[level]
    css_family = font_settings[font_choice]["css_family"]
    
    margin_in = settings["margin_in"]
    font_pt = settings["font_pt"]
    spacing_pt = settings["space_after_pt"]
    line_ht = settings["line_height"]
    
    html_content = f"""
    <div style="width: 100%; display: flex; justify-content: center; background: #f0f2f6; padding: 20px 0;">
        <div style="
            background-color: white; 
            color: black; 
            width: 8.5in; 
            min-height: 11in; 
            padding: {margin_in}in; 
            box-sizing: border-box;
            box-shadow: 0px 10px 20px rgba(0,0,0,0.15);
            font-family: {css_family};
            font-size: {font_pt}pt;
            line-height: {line_ht};
            zoom: 0.85;
            ">
            
            <div style="text-align: center; margin-bottom: 12pt;">
                <h1 style="margin: 0; font-size: {font_pt + 6}pt; font-weight: bold; letter-spacing: 0.5px; line-height: 1.0;">ALEX</h1>
                <p style="margin: 0; margin-top: 2pt; font-size: {font_pt}pt;">Boston, MA | alex@example.com | 555-0199</p>
            </div>
    """
    
    for line in resume_text.split('\n'):
        if not line.strip(): continue
        
        if line.strip().startswith('###'):
            clean_header = line.replace('###', '').strip().upper()
            html_content += f"<h3 style='margin: 0; margin-top: 12pt; margin-bottom: 2pt; font-size: {font_pt + 1}pt; font-weight: bold;'>{clean_header}</h3>"
            
        elif line.strip().startswith('**'):
            parts = [p.strip() for p in line.replace('**', '').split('|')]
            
            left_text = f"<b>{parts[0]}</b> | <b>{parts[1]}</b>" if len(parts) >= 2 else ""
            right_text = ""
            if len(parts) == 3:
                right_text = parts[2]
            elif len(parts) >= 4:
                right_text = f"{parts[2]} | {parts[3]}".replace("N/A | ", "").replace(" | N/A", "")
                
            html_content += f"<div style='display: flex; justify-content: space-between; margin-top: 6pt; margin-bottom: 2pt;'><span>{left_text}</span><span>{right_text}</span></div>"
            
        elif line.strip().startswith('-'):
            clean_bullet = line.replace('-', '', 1).strip()
            # Explicitly removed bottom margins to prevent CSS flexbox collision
            html_content += f"<ul style='margin: 0; padding-left: 24px;'><li style='margin-bottom: {spacing_pt}pt; padding-left: 4px;'>{clean_bullet}</li></ul>"
            
        else:
            html_content += f"<p style='margin: 0; margin-bottom: {spacing_pt}pt;'>{line.strip()}</p>"
            
    html_content += """
        </div>
    </div>
    """
    return html_content

# 5. STREAMLIT INTERFACE
st.title("Word Export Testing Lab 🧪")

col1, col2 = st.columns([1, 2])

with col1:
    st.markdown("### Controls")
    font_choice = st.selectbox("Select Resume Font:", options=list(font_settings.keys()))
    density_choice = st.slider("Spacing Density (1=Tight, 5=Loose):", 1, 5, 3, 1)
    
    st.divider()
    
    docx_data = create_polished_word_doc(dummy_ai_response, level=density_choice, font_choice=font_choice)
    st.download_button(
        label="💾 Download Word Document", 
        data=docx_data, 
        file_name="Polished_Test_Resume.docx", 
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )

with col2:
    st.markdown("### Live Preview (Strict Print Parity)")
    preview_html = render_live_preview(dummy_ai_response, level=density_choice, font_choice=font_choice)
    st.components.v1.html(preview_html, height=850, scrolling=True)